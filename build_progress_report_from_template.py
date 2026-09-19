from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from build_progress_report import draw_gantt, draw_pdm, draw_risk_matrix


TEMPLATE = Path("FYP Progress Report Template 2026_27.docx")
OUT = Path("FYP_Progress_Report_Zhang_Ziyi_Updated.docx")
ASSET_DIR = Path("progress_report_assets")
ASSET_DIR.mkdir(exist_ok=True)

PROJECT_TITLE = "Real-time Collaborative Music Composition Tool"
PROJECT_NO = "12"
STUDENT_ID = "P2323698"
STUDENT_NAME = "张子怡"
SUPERVISOR = "Franky Hoi Un Cheang"
SUBMISSION_DATE = "June 4, 2026"


def clear_document_body(doc):
    body = doc._body._element
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def style_table_like_template(table, header=True):
    table.style = "Table Grid"
    table.autofit = True
    for r, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10)
        if header and r == 0:
            for cell in row.cells:
                set_cell_shading(cell, "E8EEF5")
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.bold = True


def add_page_break(doc):
    doc.add_section(WD_SECTION.NEW_PAGE)


def add_body(doc, text):
    p = doc.add_paragraph(text)
    p.style = "Normal"
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Paragraph")
        set_numbering(p, num_id=2)
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Paragraph")
        set_numbering(p, num_id=1)
        p.add_run(item)


def set_numbering(paragraph, num_id, level=0):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = num_pr.find(qn("w:ilvl"))
    if ilvl is None:
        ilvl = OxmlElement("w:ilvl")
        num_pr.append(ilvl)
    ilvl.set(qn("w:val"), str(level))
    num = num_pr.find(qn("w:numId"))
    if num is None:
        num = OxmlElement("w:numId")
        num_pr.append(num)
    num.set(qn("w:val"), str(num_id))


def add_caption(doc, text):
    p = doc.add_paragraph(text)
    p.style = "Caption"
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_cover(doc):
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph("Faculty of Applied Sciences\nBachelor of Science in Computing")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("COMP4299 Final Year Project\nProgress Report\n\nAcademic Year 2026/27")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.bold = True
        run.font.size = Pt(18)
    doc.add_paragraph()
    table = doc.add_table(rows=8, cols=2)
    table.style = "Table Grid"
    rows = [
        ("Project title:", PROJECT_TITLE),
        ("Project number:", PROJECT_NO),
        ("Student ID:", STUDENT_ID),
        ("Student name:", STUDENT_NAME),
        ("Supervisor:", SUPERVISOR),
        ("Assessor:", ""),
        ("Submission Date:", SUBMISSION_DATE),
        (
            "Current project status:",
            "Working MVP completed: homepage, dashboard, demo room, Socket.io sync, local event persistence and upload filename sync.",
        ),
    ]
    for row, (k, v) in zip(table.rows, rows):
        set_cell_text(row.cells[0], k, bold=True)
        set_cell_text(row.cells[1], v)
    add_page_break(doc)


def add_front_matter(doc):
    doc.add_heading("Table of Contents", level=1)
    contents = [
        "1 Introduction",
        "1.1 Societal, User and Business Needs",
        "1.2 Objectives",
        "1.3 Ethical Consideration",
        "1.4 Summary",
        "2 Background and Related Works",
        "3 Project Management and Risk Management",
        "3.1 Project Time Management",
        "3.2 Project Risk Management",
        "3.3 Monthly Status Review",
        "4 Completed Work",
        "5 Conclusion",
        "References",
        "Appendix A. Ethics Checklist",
        "Appendix B. Declaration of the Use of Generative AI in FYP",
    ]
    for item in contents:
        add_body(doc, item)
    doc.add_heading("Table of Figures", level=1)
    for item in [
        "Figure 1: Activity List",
        "Figure 2: Precedence Diagramming Method Diagram",
        "Figure 3: Gantt Chart",
        "Figure 4: Probability impact matrix before proposed response",
    ]:
        add_body(doc, item)
    doc.add_heading("List of Tables", level=1)
    add_body(doc, "Table 1: Table of prioritized risk")
    add_body(doc, "Table 2: Current prototype function summary")
    add_page_break(doc)


def add_intro(doc):
    doc.add_heading("Introduction", level=1)
    add_body(doc, "Music production is increasingly shaped by remote work and online creative collaboration. Musicians, producers and students often need to exchange ideas without being in the same studio, but many common music creation workflows still depend on sending project files, stems or audio clips back and forth. This makes collaborative editing slow, especially when several users need to comment on timing, arrange clips, or review changes repeatedly.")
    add_body(doc, "The problem addressed by this project is that online music collaboration is difficult when users cannot view and edit the same project state at the same time. Continuous low-latency raw audio streaming is technically demanding under ordinary network conditions, so this project takes a more achievable approach: it focuses on synchronising user actions and shared project metadata in real time, while uploaded audio clips are stored and referenced by the system.")
    add_body(doc, "The project therefore aims to build a browser-based collaborative music composition tool where users create rooms, manage tracks and clips, chat, and keep the shared music workspace aligned through event-based collaboration.")
    doc.add_heading("Societal, User and Business Needs", level=2)
    add_body(doc, "The project responds to social and user needs created by remote learning, distributed creative teams and online content production. A browser-based collaborative tool can reduce the need for physical studio access and allow users in different places to work on early-stage composition ideas together.")
    add_body(doc, "Target users need a simple interface for creating rooms, uploading clips, arranging tracks, chatting and seeing other users' edits quickly. From a business and product perspective, the project has value because many creative platforms are moving towards cloud-based collaboration. A lightweight implementation can demonstrate how real-time event synchronisation may support collaborative music creation without requiring specialist audio-network infrastructure.")
    add_body(doc, "The main ethical and legal needs are privacy, account data protection and appropriate handling of uploaded audio. The system should avoid collecting unnecessary personal data, protect user account information, and remind users that they should upload only audio they have the right to use.")
    doc.add_heading("Objectives", level=2)
    add_numbered(doc, [
        "Build a web-based prototype that allows users to create or join a collaboration room.",
        "Provide a polished homepage, dashboard and room workspace suitable for a Final Year Project prototype demonstration.",
        "Provide basic audio clip upload metadata and listing functions so that users can share short audio materials within a room.",
        "Implement synchronisation of user actions such as uploading clip metadata, moving clips, marking sections, adding notes and syncing edits using a room-event update channel.",
        "Design a clear and usable interface for room management, track display, chat and simple clip arrangement.",
        "Test the prototype with representative scenarios and evaluate whether shared project state remains consistent across multiple clients.",
    ])
    doc.add_heading("Ethical Consideration", level=2)
    add_body(doc, "The ethics checklist is included in Appendix A. The project does not require human-subject experiments at this stage, but it may store basic user account information and uploaded audio clips in the final version. The design will therefore minimise personal data collection and use the data only for project demonstration purposes.")
    doc.add_heading("Summary", level=2)
    add_body(doc, "This report first introduces the project problem and objectives, then reviews related work on networked music collaboration and low-latency audio systems. It then presents project management plans, risk management, the monthly status review, and the work completed so far. Since a first working web prototype has now been developed, the report focuses on the implemented foundation, current limitations, and the remaining work needed before the final submission.")


def add_background(doc):
    doc.add_heading("Background and Related Works", level=1)
    add_body(doc, "Networked music collaboration has been discussed for many years, but the technical requirements vary depending on whether the system aims to stream live audio performance or support asynchronous or semi-synchronous composition. Oliveros [1] highlights that networked music can expand creative practice, but also introduces technical and communication challenges.")
    add_body(doc, "Systems such as JackTrip [2] show the importance of low-latency audio transmission for real-time performance. However, such systems often require careful setup, strong network conditions and specialist knowledge. For a general web-based student project, building a full low-latency streaming engine would create high technical risk.")
    add_body(doc, "More recent work on Audio over OSC (AOO) [3] demonstrates message-based communication for audio and metadata, which supports the idea that not all collaboration has to depend on continuous raw audio streaming. Research evaluating networked music performance over LEO satellite internet [4] further confirms that latency, jitter and packet loss remain difficult challenges even when advanced network infrastructure is used.")
    add_body(doc, "Based on these related works, this project narrows its scope. Instead of attempting professional live-performance streaming, it implements event-based synchronisation for a shared music composition workspace. The expected contribution is a practical prototype that demonstrates collaborative project-state synchronisation in a music-oriented interface.")


def add_management(doc):
    doc.add_heading("Project Management and Risk Management", level=1)
    add_body(doc, "This chapter explains how the project is managed after the progress report stage. A working MVP has been completed, including a Socket.io-based event synchronisation layer and local JSON event persistence. The remaining schedule now prioritises improving real audio-clip handling, strengthening multi-client testing, and preparing the final evaluation.")
    doc.add_heading("Project Time Management", level=2)
    activities = [
        ("A1", "Confirm project scope, tools and development environment", "2 days", "-"),
        ("A2", "Requirement analysis and user scenario definition", "3 days", "A1"),
        ("A3", "Background research and related work review", "5 days", "A1"),
        ("A4", "System architecture and data model design", "3 days", "A2"),
        ("A5", "User interface prototype and room workflow design", "3 days", "A4"),
        ("A6", "Local prototype server and room API implementation", "5 days", "A4"),
        ("A7", "Frontend homepage, dashboard and room interface implementation", "7 days", "A5"),
        ("A8", "Event synchronisation prototype using room update events", "4 days", "A6, A7"),
        ("A9", "Audio upload metadata and clip movement prototype", "4 days", "A6, A7"),
        ("A10", "Persistence, Socket.io refinement and multi-client testing", "7 days", "A8, A9"),
        ("A11", "Final report writing, evaluation and presentation preparation", "10 days", "A10"),
    ]
    table = doc.add_table(rows=1, cols=4)
    for i, h in enumerate(["Activity", "Description", "Duration", "Predecessor"]):
        set_cell_text(table.rows[0].cells[i], h, bold=True)
    for row in activities:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            set_cell_text(cells[i], val)
    style_table_like_template(table)
    add_caption(doc, "Figure 1: Activity List")
    pdm = ASSET_DIR / "pdm.png"
    gantt = ASSET_DIR / "gantt.png"
    draw_pdm(pdm)
    draw_gantt(gantt)
    doc.add_picture(str(pdm), width=Inches(6.1))
    add_caption(doc, "Figure 2: Precedence Diagramming Method Diagram")
    doc.add_picture(str(gantt), width=Inches(6.1))
    add_caption(doc, "Figure 3: Gantt Chart")
    doc.add_heading("Project Risk Management", level=2)
    risks = [
        ("1", "R1: Real-time synchronisation is more complex than expected", "High", "High", "Reduce scope to event synchronisation first; test with two clients before adding optional features."),
        ("2", "R2: Development time is limited and advanced functions remain incomplete", "High", "Medium", "Keep the prototype focused on room creation, clip management, track editing and event synchronisation; defer advanced DAW editing and live audio streaming."),
        ("3", "R3: Large audio files may slow down upload and storage", "Medium", "High", "Restrict file type and size; store metadata separately; use short audio clips for prototype testing."),
        ("4", "R4: Full-stack and real-time networking skill gaps may slow progress", "Medium", "Medium", "Use established frameworks such as Next.js, TypeScript, TailwindCSS and Socket.io; build in small increments."),
    ]
    table = doc.add_table(rows=1, cols=5)
    for i, h in enumerate(["Priority", "Risk Identifier and Description", "Probability", "Impact", "Response"]):
        set_cell_text(table.rows[0].cells[i], h, bold=True)
    for row in risks:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            set_cell_text(cells[i], val)
    style_table_like_template(table)
    add_caption(doc, "Table 1: Table of prioritized risk")
    matrix = ASSET_DIR / "risk_matrix.png"
    draw_risk_matrix(matrix)
    doc.add_picture(str(matrix), width=Inches(5.5))
    add_caption(doc, "Figure 4: Probability impact matrix before proposed response")
    doc.add_heading("Monthly Status Review", level=2)
    table = doc.add_table(rows=1, cols=4)
    for i, h in enumerate(["Period", "Work completed", "Current issue", "Plan for next stage"]):
        set_cell_text(table.rows[0].cells[i], h, bold=True)
    rows = [
        ("March 2026", "Project proposal completed; project problem, objectives, risks and preliminary references identified.", "The implementation scope was still broad.", "Narrow the project to event-based collaboration instead of full live audio streaming."),
        ("April 2026", "Related work reviewed; core technical direction confirmed: room-based web platform, audio clip upload and state synchronisation.", "Need to convert proposal into an implementation plan.", "Prepare architecture, data model and UI workflow."),
        (
            "May 2026",
            "Working MVP developed, including a visual homepage, dashboard, demo room, waveform-based editing UI, upload audio button, event stream, room-scoped notes, Socket.io real-time event synchronisation and local JSON event persistence.",
            "The prototype is now functional for demonstration, but real audio decoding, file storage and production database support are still limited.",
            "Stabilise the prototype, document testing evidence, and prepare the progress report demonstration.",
        ),
        (
            "June 2026",
            "Planned submission of progress report and prototype demonstration.",
            "Time is the main constraint; optional features must not distract from the core collaboration loop.",
            "Complete multi-client testing, refine audio upload handling, prepare screenshots and write the final report plan.",
        ),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            set_cell_text(cells[i], val)
    style_table_like_template(table)


def add_completed_work(doc):
    doc.add_heading("Completed Work", level=1)
    add_body(doc, "At the progress report stage, a working MVP has been completed. The current implementation is not intended to be a full digital audio workstation or a low-latency live performance streaming system. Instead, it demonstrates the main project direction: collaborative music editing through synchronised user actions and shared room state.")
    add_bullets(doc, [
        "A polished homepage has been redesigned as a strong visual entry point, using floating music objects, lightweight motion and a dark premium visual system instead of placing the editor screenshot directly on the landing page.",
        "A dashboard page has been implemented where users can view mock music rooms and navigate into the demo collaboration workspace.",
        "A demo collaboration room interface has been implemented with a top bar, online collaborator avatars, track arrangement sidebar, waveform editor, playback controls, upload audio button, event activity stream and room-scoped notes panel.",
        "The waveform editor has been simplified so that red, yellow, blue and green waveform segments represent different editing actions without heavy selection boxes. The current actions are Move clip, Mark section, Sync edit, Add note and Extend clip.",
        "The upload audio button allows users to select a local audio file, display the chosen filename in the interface and synchronise the filename to another connected browser window.",
        "A custom Next.js server has been implemented with Socket.io. When one browser window triggers a room event such as Move clip, another connected browser window receives the event in real time.",
        "Simple backend API routes have been implemented for creating/reading rooms, reading room events, posting new events and updating room metadata.",
        "Event history is persisted to a local JSON file, so the demo room can reload previously generated events instead of losing all state immediately after refresh.",
        "The visual interface has been refined with a consistent red, yellow, blue and green accent system, restrained glass surfaces, responsive layout behaviour and improved readability.",
    ])
    doc.add_heading("Current Prototype Functions", level=2)
    functions = [
        ("Homepage", "Provides a visually strong landing page with floating music objects, lightweight animation and links to the dashboard and demo room."),
        ("Dashboard", "Shows mock music rooms and provides navigation into the demo collaboration workspace."),
        ("Collaboration room", "Provides the simplified collaborative music editor layout with tracks, waveform editor, controls, upload UI, event stream, notes panel and online collaborator status."),
        ("Track display", "Shows four arrangement lanes with clip counts and mute/status icons to represent shared project state."),
        ("Waveform editor", "Shows coloured waveform segments for Move clip, Mark section, Sync edit and Extend clip. Clicking a toolbar action or waveform segment creates a room event."),
        ("Upload workflow", "Allows users to select a local audio file, display the filename, update the source clip label and broadcast the upload metadata to other connected clients."),
        ("Room notes", "Displays room-scoped creative notes and supports Add note as an event action. Full typed chat sending is planned for the next stage."),
        ("Socket.io sync", "Uses Socket.io to broadcast room events and upload metadata between browser windows in real time."),
        ("Persistence", "Stores room metadata and event history in a local JSON file through backend API routes."),
    ]
    table = doc.add_table(rows=1, cols=3)
    for i, h in enumerate(["Area", "Implemented function", "Status"]):
        set_cell_text(table.rows[0].cells[i], h, bold=True)
    for area, function in functions:
        cells = table.add_row().cells
        set_cell_text(cells[0], area)
        set_cell_text(cells[1], function)
        set_cell_text(cells[2], "Prototype completed")
    style_table_like_template(table)
    add_caption(doc, "Table 2: Current prototype function summary")
    doc.add_heading("Limitations and Remaining Work", level=2)
    add_bullets(doc, [
        "The current persistence layer uses a local JSON file for prototype demonstration. A production-ready database such as MongoDB should be added before the final version.",
        "The prototype does not yet include user registration, authentication or role-based room access.",
        "The upload function currently focuses on file selection, filename display and metadata synchronisation. Real file storage, playback and waveform decoding should be improved in the next stage.",
        "The room notes panel is currently partly visual. Full typed chat sending and persistence should be implemented or clearly scoped for the final demo.",
        "Track editing is currently represented by static arrangement lanes. Adding, renaming and muting tracks should be connected to real room events if time allows.",
        "The Socket.io event layer works in local browser-window testing, but more systematic multi-client testing is required.",
        "The final evaluation should include test scenarios, screenshots, browser-window synchronisation evidence and a clear explanation of why event-based collaboration is used instead of raw audio streaming.",
    ])


def add_conclusion_refs_appendices(doc):
    doc.add_heading("Conclusion", level=1)
    add_body(doc, "This progress report shows that the project has moved beyond proposal and planning into a working MVP. The current website demonstrates the central idea of the project: users work in a shared music room, view and manage music-related objects, upload clip metadata, and synchronise editing actions rather than streaming live raw audio.")
    add_body(doc, "For the next stage, I will avoid adding optional features until the core collaboration loop is reliable and well tested. Since Socket.io synchronisation and local persistence are now implemented, the priorities are to improve real audio clip handling, strengthen typed room chat and track editing, test the prototype with multiple browser clients, and prepare clear evidence for the final report and presentation.")
    doc.add_heading("References", level=1)
    refs = [
        "[1] Pauline Oliveros. Networked music: Low and high tech. Contemporary Music Review, 28(4-5):433-435, 2009.",
        "[2] Juan-Pablo Caceres and Chris Chafe. JackTrip: Under the hood of an engine for network audio. Journal of New Music Research, 39(3):183-187, 2010.",
        "[3] Ivica Ico Ressi. AOO: Low-latency peer-to-peer audio streaming and messaging. https://www.soundingfuture.com/en/article/aoo-low-latency-peer-peer-audio-streaming-and-messaging, 2024 [accessed Mar. 18, 2026].",
        "[4] Luca Borgianni, Daniele Adami, Marco Bosi, Stefano Giordano and Chris Chafe. A comprehensive evaluation of networked music performance using LEO satellite internet: The Starlink use case. IEEE Transactions on Network and Service Management, 22(5):3947-3963, Oct. 2025.",
    ]
    for ref in refs:
        add_body(doc, ref)
    doc.add_heading("Appendix A. Ethics Checklist", level=1)
    meta = doc.add_table(rows=5, cols=2)
    for row, (k, v) in zip(meta.rows, [
        ("Name of student:", STUDENT_NAME),
        ("Student number:", STUDENT_ID),
        ("Programme:", "BSc in Computing"),
        ("Module:", "Final Year Project"),
        ("Title of project:", PROJECT_TITLE),
    ]):
        set_cell_text(row.cells[0], k, bold=True)
        set_cell_text(row.cells[1], v)
    style_table_like_template(meta, header=False)
    add_body(doc, "Project summary: This project develops a browser-based collaborative music composition prototype. Users can enter a shared room, view track lanes, select a local audio file, trigger waveform editing actions, display room-scoped notes and synchronise editing events through a real-time room-event channel. The project is intended as a technical prototype and does not require formal user experiments at the progress report stage.")
    qtable = doc.add_table(rows=1, cols=5)
    for i, h in enumerate(["No.", "Question", "Yes", "No", "N/A"]):
        set_cell_text(qtable.rows[0].cells[i], h, bold=True)
    questions = [
        ("1", "Does your project include human participants?", "", "", "✓"),
        ("2", "Will any participants be from vulnerable groups?", "", "", "✓"),
        ("3", "Will you tell participants that participation is voluntary?", "", "", "✓"),
        ("4", "Will you obtain written consent for participation?", "", "", "✓"),
        ("5", "Will participants be told that they may withdraw at any time?", "", "", "✓"),
        ("6", "Is there any realistic risk of physical or psychological distress?", "", "✓", ""),
        ("7", "Will any non-anonymised and/or personalised data be generated and/or stored?", "✓", "", ""),
        ("8", "Will you access documents containing sensitive data about living individuals?", "", "✓", ""),
        ("9", "Will you be exposed to risks greater than normal study/working life?", "", "✓", ""),
        ("10", "Will you be exposed to highly addictive or illegal activities?", "", "✓", ""),
    ]
    for row in questions:
        cells = qtable.add_row().cells
        for i, val in enumerate(row):
            set_cell_text(cells[i], val)
    style_table_like_template(qtable)
    add_body(doc, "Potential ethics issue handling: The prototype may store basic account details and uploaded audio clips. The system design will minimise personal data, use the information only for project demonstration, and avoid collecting sensitive information. Users should upload only audio material they are allowed to use.")
    add_body(doc, "Declaration: I have read the instructions carefully and reported honestly the potential ethics issues about the project.")
    doc.add_heading("Appendix B. Declaration of the Use of Generative AI in FYP", level=1)
    add_body(doc, "☒ I have used generative AI in this assessment work")
    add_body(doc, "☐ I did not use generative AI in this assessment work")
    add_body(doc, "Generative AI tool used: OpenAI ChatGPT / Codex.")
    table = doc.add_table(rows=1, cols=3)
    for i, h in enumerate(["Type of usage", "GenAI assistance received", "Student action on generated content"]):
        set_cell_text(table.rows[0].cells[i], h, bold=True)
    rows = [
        ("Brainstorming and idea generation", "Helped organise the project problem, objectives and scope.", "Reviewed and kept content consistent with the approved proposal."),
        ("Literature review and background research", "Helped summarise related work already listed in the proposal.", "Checked that references match the project topic."),
        ("Programming support", "Helped scaffold and refine the first working website prototype, including UI structure, local preview server logic and event-based interaction flow.", "Reviewed the generated code, tested the website locally and kept the implementation aligned with the approved project scope."),
        ("Visualization and diagram production", "Helped generate activity, PDM, Gantt and risk-management visuals.", "Reviewed diagrams for consistency with the schedule."),
        ("Text Rewriting and Drafting", "Helped draft report sections in academic English.", "Reviewed, edited and accepted responsibility for the final wording."),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            set_cell_text(cells[i], val)
    style_table_like_template(table)
    add_body(doc, "Student signature: ____________________________    Date: " + SUBMISSION_DATE)


def build():
    doc = Document(TEMPLATE)
    clear_document_body(doc)
    add_cover(doc)
    add_front_matter(doc)
    add_intro(doc)
    add_background(doc)
    add_management(doc)
    add_completed_work(doc)
    add_conclusion_refs_appendices(doc)
    doc.save(OUT)


if __name__ == "__main__":
    build()
